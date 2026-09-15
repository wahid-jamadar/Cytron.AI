"""
modules/documentation_generator/exporters.py
─────────────────────────────────────────────
Strategy pattern and factory for exporting generated documentation to various output formats:
Markdown, HTML, PDF, DOCX, Plain Text, JSON, YAML, XML, Draw.io, and PowerPoint (PPTX).
"""

import abc
import json
import os
import yaml
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List

# Exporter dependency guards
try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    docx = None

try:
    import pptx
    from pptx.util import Inches as PtInches, Pt as PtSize
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor as PptRGBColor
except ImportError:
    pptx = None

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
except ImportError:
    letter = None

from modules.documentation_generator.parser import parse_markdown_to_blocks
from modules.documentation_generator.drawio_parser import generate_drawio_from_markdown


class BaseExporter(abc.ABC):
    @abc.abstractmethod
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        """Exports the generated markdown to the targeted output path."""
        pass


class MarkdownExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(content_markdown)


class HTMLExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        import markdown
        
        # Convert MD to HTML
        html_content = markdown.markdown(
            content_markdown, 
            extensions=['tables', 'fenced_code', 'toc']
        )
        
        title = metadata.get("title", "Generated Documentation")
        doc_type = metadata.get("type", "Documentation")
        tone = metadata.get("tone", "Professional")
        date_str = metadata.get("date", datetime.utcnow().strftime("%Y-%m-%d"))

        # Premium styled HTML template
        full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap" rel="stylesheet">
    <style>
        :root {{
            --bg: #07080d;
            --surface: #0f1117;
            --border: rgba(255, 255, 255, 0.08);
            --text: #e8eaf0;
            --text-muted: #9ca3af;
            --accent: #7c3aed;
            --accent-glow: rgba(124, 58, 237, 0.25);
            --font: 'Inter', system-ui, sans-serif;
            --mono: 'JetBrains Mono', monospace;
        }}
        body {{
            font-family: var(--font);
            background-color: var(--bg);
            color: var(--text);
            line-height: 1.7;
            padding: 2rem 4vw 4rem;
            margin: 0;
        }}
        .container {{
            max-width: 900px;
            margin: 0 auto;
            background-color: var(--surface);
            border: 1px solid var(--border);
            border-radius: 12px;
            padding: 3rem;
            box-shadow: 0 10px 30px rgba(0,0,0,0.5);
        }}
        header {{
            border-bottom: 1px solid var(--border);
            padding-bottom: 2rem;
            margin-bottom: 2.5rem;
        }}
        h1 {{
            font-size: 2.5rem;
            font-weight: 800;
            margin-top: 0;
            margin-bottom: 0.5rem;
            background: linear-gradient(135deg, #fff 0%, #a855f7 100%);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
        }}
        .metadata {{
            font-size: 0.9rem;
            color: var(--text-muted);
            display: flex;
            gap: 1.5rem;
            margin-bottom: 1rem;
        }}
        .metadata span strong {{
            color: var(--text);
        }}
        h2 {{ font-size: 1.75rem; border-bottom: 1px solid var(--border); padding-bottom: 0.5rem; margin-top: 2rem; }}
        h3 {{ font-size: 1.35rem; margin-top: 1.5rem; }}
        p {{ margin-bottom: 1.25rem; }}
        ul, ol {{ margin-bottom: 1.25rem; padding-left: 1.5rem; }}
        li {{ margin-bottom: 0.5rem; }}
        code {{
            font-family: var(--mono);
            font-size: 0.9em;
            background: rgba(255,255,255,0.06);
            padding: 0.2rem 0.4rem;
            border-radius: 4px;
            color: #f472b6;
        }}
        pre {{
            background: #161b26;
            border: 1px solid var(--border);
            border-radius: 8px;
            padding: 1.25rem;
            overflow-x: auto;
            margin-bottom: 1.5rem;
        }}
        pre code {{
            background: none;
            padding: 0;
            color: inherit;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 1.5rem;
            font-size: 0.95rem;
        }}
        th, td {{
            padding: 0.75rem 1rem;
            border: 1px solid var(--border);
            text-align: left;
        }}
        th {{
            background-color: rgba(124, 58, 237, 0.15);
            font-weight: 600;
            color: #fff;
        }}
        tr:nth-child(even) {{
            background-color: rgba(255,255,255,0.02);
        }}
    </style>
</head>
<body>
    <div class="container">
        <header>
            <h1>{title}</h1>
            <div class="metadata">
                <span><strong>Type:</strong> {doc_type}</span>
                <span><strong>Tone:</strong> {tone}</span>
                <span><strong>Date:</strong> {date_str}</span>
            </div>
        </header>
        <main>
            {html_content}
        </main>
    </div>
</body>
</html>
"""
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(full_html)


class PlainTextExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        blocks = parse_markdown_to_blocks(content_markdown)
        lines: List[str] = []

        title = metadata.get("title", "DOCUMENTATION").upper()
        doc_type = metadata.get("type", "General").upper()
        date_str = metadata.get("date", datetime.utcnow().strftime("%Y-%m-%d"))

        # Add Title Banner
        lines.append("=" * 80)
        lines.append(f" {title}")
        lines.append(f" Type: {doc_type} | Date: {date_str}")
        lines.append("=" * 80)
        lines.append("")

        for block in blocks:
            b_type = block["type"]
            if b_type == "heading":
                level = block["level"]
                text = block["text"]
                if level == 1:
                    lines.append("")
                    lines.append(text.upper())
                    lines.append("=" * len(text))
                    lines.append("")
                elif level == 2:
                    lines.append("")
                    lines.append(text)
                    lines.append("-" * len(text))
                    lines.append("")
                else:
                    lines.append("")
                    lines.append(f"{'  ' * (level - 2)}* {text}")
                    lines.append("")
            elif b_type == "paragraph":
                # Basic wrapping
                text = block["text"]
                words = text.split()
                line = []
                for word in words:
                    if len(" ".join(line + [word])) > 76:
                        lines.append(" ".join(line))
                        line = [word]
                    else:
                        line.append(word)
                if line:
                    lines.append(" ".join(line))
                lines.append("")
            elif b_type == "list":
                ordered = block["ordered"]
                for idx, item in enumerate(block["items"]):
                    prefix = f"  {idx + 1}. " if ordered else "  * "
                    lines.append(f"{prefix}{item}")
                lines.append("")
            elif b_type == "code_block":
                lang = block["language"] or "Text"
                lines.append(f"--- CODE BLOCK ({lang}) ---")
                for c_line in block["code"].splitlines():
                    lines.append(f"  {c_line}")
                lines.append("-" * 40)
                lines.append("")
            elif b_type == "table":
                headers = block["headers"]
                rows = block["rows"]
                
                # Determine column widths
                widths = [len(h) for h in headers]
                for row in rows:
                    for col_idx, cell in enumerate(row):
                        if col_idx < len(widths):
                            widths[col_idx] = max(widths[col_idx], len(str(cell)))
                
                # Format grid
                sep = "+" + "+".join(["-" * (w + 2) for w in widths]) + "+"
                lines.append(sep)
                
                header_line = "|" + "|".join([f" {headers[c_idx].ljust(widths[c_idx])} " for c_idx in range(len(headers))]) + "|"
                lines.append(header_line)
                lines.append(sep)
                
                for row in rows:
                    cells_str = []
                    for col_idx, cell in enumerate(row):
                        if col_idx < len(widths):
                            cells_str.append(f" {str(cell).ljust(widths[col_idx])} ")
                        else:
                            cells_str.append(" ")
                    lines.append("|" + "|".join(cells_str) + "|")
                lines.append(sep)
                lines.append("")

        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))


class JSONExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        blocks = parse_markdown_to_blocks(content_markdown)
        data = {
            "title": metadata.get("title", "Documentation"),
            "metadata": {
                "type": metadata.get("type", "Unknown"),
                "tone": metadata.get("tone", "Professional"),
                "generated_at": datetime.utcnow().isoformat(),
            },
            "blocks": blocks
        }
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)


class YAMLExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        blocks = parse_markdown_to_blocks(content_markdown)
        data = {
            "title": metadata.get("title", "Documentation"),
            "metadata": {
                "type": metadata.get("type", "Unknown"),
                "tone": metadata.get("tone", "Professional"),
                "generated_at": datetime.utcnow().isoformat(),
            },
            "blocks": blocks
        }
        with open(output_path, "w", encoding="utf-8") as f:
            yaml.dump(data, f, default_flow_style=False, sort_keys=False)


class XMLExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        blocks = parse_markdown_to_blocks(content_markdown)
        
        root = ET.Element("document", {
            "title": metadata.get("title", "Documentation"),
            "type": metadata.get("type", "Unknown"),
            "tone": metadata.get("tone", "Professional"),
            "generated_at": datetime.utcnow().isoformat()
        })
        
        for idx, block in enumerate(blocks):
            b_type = block["type"]
            elem = ET.SubElement(root, "block", {"id": str(idx), "type": b_type})
            
            if b_type == "heading":
                elem.set("level", str(block["level"]))
                elem.text = block["text"]
            elif b_type == "paragraph":
                elem.text = block["text"]
            elif b_type == "code_block":
                elem.set("language", block["language"])
                code_elem = ET.SubElement(elem, "code")
                code_elem.text = block["code"]
            elif b_type == "list":
                elem.set("ordered", str(block["ordered"]).lower())
                for item in block["items"]:
                    item_elem = ET.SubElement(elem, "item")
                    item_elem.text = item
            elif b_type == "table":
                headers_elem = ET.SubElement(elem, "headers")
                for h in block["headers"]:
                    h_elem = ET.SubElement(headers_elem, "header")
                    h_elem.text = h
                
                rows_elem = ET.SubElement(elem, "rows")
                for r in block["rows"]:
                    row_elem = ET.SubElement(rows_elem, "row")
                    for cell in r:
                        cell_elem = ET.SubElement(row_elem, "cell")
                        cell_elem.text = str(cell)
                        
        tree = ET.ElementTree(root)
        tree.write(output_path, encoding="utf-8", xml_declaration=True)


class DrawioExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        title = metadata.get("title", "Architecture Diagram")
        drawio_xml = generate_drawio_from_markdown(content_markdown, title)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(drawio_xml)


class DOCXExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        if not docx:
            raise RuntimeError("python-docx is not installed in the Python environment")

        doc = docx.Document()

        # Page Setup
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Style configurations
        style_normal = doc.styles['Normal']
        style_normal.font.name = 'Arial'
        style_normal.font.size = Pt(11)
        style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        title_text = metadata.get("title", "Project Documentation")
        doc_type = metadata.get("type", "Technical Document")
        tone = metadata.get("tone", "Professional")
        date_str = metadata.get("date", datetime.utcnow().strftime("%Y-%m-%d"))

        # Add Title Page
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(f"\n\n\n\n\n\n{title_text}\n")
        title_run.font.size = Pt(28)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x7c, 0x3a, 0xed) # Accent purple

        subtitle_p = doc.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle_p.add_run(f"Type: {doc_type} | Tone: {tone}\nDate: {date_str}\n")
        subtitle_run.font.size = Pt(14)
        subtitle_run.font.italic = True
        subtitle_run.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)

        doc.add_page_break()

        # Parse and populate body
        blocks = parse_markdown_to_blocks(content_markdown)
        for block in blocks:
            b_type = block["type"]
            
            if b_type == "heading":
                level = block["level"]
                # python-docx headings levels are 0-9
                h = doc.add_heading(block["text"], level=level)
                # Set purple heading color for levels 1 & 2
                if level <= 2:
                    for run in h.runs:
                        run.font.color.rgb = RGBColor(0x7c, 0x3a, 0xed)
            
            elif b_type == "paragraph":
                doc.add_paragraph(block["text"])
                
            elif b_type == "list":
                ordered = block["ordered"]
                style_name = 'List Number' if ordered else 'List Bullet'
                for item in block["items"]:
                    doc.add_paragraph(item, style=style_name)
                    
            elif b_type == "code_block":
                # Create a single cell table to box the code block
                table = doc.add_table(rows=1, cols=1)
                table.style = 'Light Shading Accent 1'
                cell = table.cell(0, 0)
                
                # Shading color background: light grey
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'clear')
                shading.set(qn('w:color'), 'auto')
                shading.set(qn('w:fill'), 'F3F4F6') # Light grey
                cell._tc.get_or_add_tcPr().append(shading)
                
                p = cell.paragraphs[0]
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.right_indent = Inches(0.2)
                
                code_run = p.add_run(block["code"])
                code_run.font.name = 'Courier New'
                code_run.font.size = Pt(9.5)
                code_run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
                
                # Blank separator paragraph after table
                doc.add_paragraph()
                
            elif b_type == "table":
                headers = block["headers"]
                rows = block["rows"]
                
                doc_table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                doc_table.style = 'Light Shading Accent 1'
                
                # Format headers
                hdr_cells = doc_table.rows[0].cells
                for idx, text in enumerate(headers):
                    hdr_cells[idx].text = text
                    p = hdr_cells[idx].paragraphs[0]
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        
                    # Set purple background shading for headers
                    tcPr = hdr_cells[idx]._tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), '7C3AED') # Accent Purple
                    tcPr.append(shd)
                
                # Populate rows
                for r_idx, row_data in enumerate(rows):
                    row_cells = doc_table.rows[r_idx + 1].cells
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < len(row_cells):
                            row_cells[c_idx].text = str(cell_value)
                            
                # Blank separator paragraph after table
                doc.add_paragraph()

        doc.save(str(output_path))


class PowerPointExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        if not pptx:
            raise RuntimeError("python-pptx is not installed in the Python environment")

        prs = pptx.Presentation()
        
        # 1. Title Slide (Layout 0)
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        
        title_box = slide.shapes.title
        subtitle_box = slide.placeholders[1]
        
        title_text = metadata.get("title", "Project Documentation")
        doc_type = metadata.get("type", "Technical Specifications")
        tone = metadata.get("tone", "Professional")
        date_str = metadata.get("date", datetime.utcnow().strftime("%Y-%m-%d"))
        
        title_box.text = title_text
        subtitle_box.text = f"Type: {doc_type}\nTone: {tone}\nDate: {date_str}\nCreated by Autonomous Agent System"
        
        # Parse blocks
        blocks = parse_markdown_to_blocks(content_markdown)
        
        # 2. Sequential slides based on Headings (Layout 1: Title & Content)
        current_slide = None
        current_text_frame = None
        
        for block in blocks:
            b_type = block["type"]
            
            if b_type == "heading" and block["level"] <= 2:
                # Create a new slide for Heading 1 & Heading 2
                slide_layout = prs.slide_layouts[1]
                current_slide = prs.slides.add_slide(slide_layout)
                
                # Title
                title_shape = current_slide.shapes.title
                title_shape.text = block["text"]
                
                # Content placeholder
                body_shape = current_slide.placeholders[1]
                current_text_frame = body_shape.text_frame
                current_text_frame.clear() # clear default bullets
                
            elif current_slide is not None and current_text_frame is not None:
                if b_type == "paragraph":
                    p = current_text_frame.add_paragraph()
                    p.text = block["text"]
                    p.level = 0
                    p.space_after = PtSize(10)
                    
                elif b_type == "list":
                    for item in block["items"]:
                        p = current_text_frame.add_paragraph()
                        p.text = item
                        p.level = 1
                        p.space_after = PtSize(5)
                        
                elif b_type == "code_block":
                    p = current_text_frame.add_paragraph()
                    p.text = f"Code snippet ({block['language'] or 'Text'}):"
                    p.level = 0
                    p.font.bold = True
                    
                    p_code = current_text_frame.add_paragraph()
                    # Keep only first 5 lines of code for slide space
                    lines = block["code"].splitlines()[:5]
                    if len(block["code"].splitlines()) > 5:
                        lines.append("... [truncated] ...")
                    p_code.text = "\n".join(lines)
                    p_code.level = 1
                    p_code.font.name = "Courier New"
                    p_code.font.size = PtSize(12)
                    p_code.space_after = PtSize(10)
                    
                elif b_type == "table":
                    headers = block["headers"]
                    rows = block["rows"][:5] # Limit table size for slide layout
                    
                    # Create Title Only slide for Tables to avoid overlap
                    table_layout = prs.slide_layouts[5] # Title Only
                    table_slide = prs.slides.add_slide(table_layout)
                    table_slide.shapes.title.text = f"{current_slide.shapes.title.text} - Data Table"
                    
                    # Table Dimensions
                    left = PtInches(1.0)
                    top = PtInches(2.0)
                    width = PtInches(8.0)
                    height = PtInches(0.8 + 0.4 * len(rows))
                    
                    table_shape = table_slide.shapes.add_table(
                        rows=len(rows) + 1, 
                        cols=len(headers), 
                        left=left, 
                        top=top, 
                        width=width, 
                        height=height
                    )
                    table = table_shape.table
                    
                    # Set column headers
                    for c_idx, h_text in enumerate(headers):
                        cell = table.cell(0, c_idx)
                        cell.text = h_text
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = PptRGBColor(124, 58, 237) # Purple header
                    
                    # Populate data
                    for r_idx, row_data in enumerate(rows):
                        for c_idx, val in enumerate(row_data):
                            if c_idx < len(headers):
                                table.cell(r_idx + 1, c_idx).text = str(val)

        # 3. Final Conclusion Slide
        slide_layout = prs.slide_layouts[1]
        final_slide = prs.slides.add_slide(slide_layout)
        final_slide.shapes.title.text = "Conclusion & Next Steps"
        tf = final_slide.placeholders[1].text_frame
        tf.clear()
        
        p = tf.add_paragraph()
        p.text = "Documentation summary completed successfully."
        p.level = 0
        p.space_after = PtSize(14)
        
        p2 = tf.add_paragraph()
        p2.text = "Review completed sections in the generated documents."
        p2.level = 1
        
        prs.save(str(output_path))


class PDFExporter(BaseExporter):
    def export(self, content_markdown: str, metadata: Dict[str, Any], output_path: Path) -> None:
        if not letter:
            raise RuntimeError("reportlab is not installed in the Python environment")

        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
        )

        styles = getSampleStyleSheet()
        
        # Define Custom Styles
        title_style = ParagraphStyle(
            name='DocTitle',
            parent=styles['Title'],
            fontName='Helvetica-Bold',
            fontSize=28,
            leading=34,
            textColor=colors.HexColor('#7C3AED'),
            spaceAfter=20
        )
        
        subtitle_style = ParagraphStyle(
            name='DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#6B7280'),
            alignment=1, # Centered
            spaceAfter=200
        )
        
        h1_style = ParagraphStyle(
            name='DocH1',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#7C3AED'),
            spaceBefore=15,
            spaceAfter=10,
            keepWithNext=True
        )

        h2_style = ParagraphStyle(
            name='DocH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#7C3AED'),
            spaceBefore=12,
            spaceAfter=8,
            keepWithNext=True
        )

        body_style = ParagraphStyle(
            name='DocBody',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=10,
            leading=15,
            textColor=colors.HexColor('#222222'),
            spaceAfter=10
        )

        code_style = ParagraphStyle(
            name='DocCode',
            parent=styles['Normal'],
            fontName='Courier',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#111827'),
            spaceAfter=10
        )

        story = []

        title_text = metadata.get("title", "Project Documentation")
        doc_type = metadata.get("type", "Technical Specifications")
        tone = metadata.get("tone", "Professional")
        date_str = metadata.get("date", datetime.utcnow().strftime("%Y-%m-%d"))

        # Title Page
        story.append(Spacer(1, 150))
        story.append(Paragraph(title_text, title_style))
        story.append(Paragraph(f"Type: {doc_type} | Tone: {tone}<br/>Date: {date_str}<br/>Created by Cytron.AI system", subtitle_style))
        story.append(PageBreak())

        # Body parsing
        blocks = parse_markdown_to_blocks(content_markdown)
        for block in blocks:
            b_type = block["type"]
            
            if b_type == "heading":
                level = block["level"]
                h_style = h1_style if level == 1 else h2_style
                story.append(Paragraph(block["text"], h_style))
                
            elif b_type == "paragraph":
                story.append(Paragraph(block["text"], body_style))
                
            elif b_type == "list":
                ordered = block["ordered"]
                for idx, item in enumerate(block["items"]):
                    bullet = f"{idx + 1}. " if ordered else "&bull; "
                    story.append(Paragraph(f"{bullet}{item}", body_style))
                story.append(Spacer(1, 5))
                
            elif b_type == "code_block":
                # Escape code text for HTML formatting in reportlab Paragraphs
                escaped_code = block["code"].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                code_p = Paragraph(f"<pre>{escaped_code}</pre>", code_style)
                
                # Place in Single-cell Table for Border Box
                t = Table([[code_p]], colWidths=[500])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F3F4F6')),
                    ('PADDING', (0,0), (-1,-1), 10),
                    ('LEFTPADDING', (0,0), (-1,-1), 15),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                    ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
                    ('LINEABOVE', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
                    ('LINELEFT', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
                    ('LINERIGHT', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))
                
            elif b_type == "table":
                headers = block["headers"]
                rows = block["rows"]
                
                # Wrap cell data in Paragraphs to auto-wrap text inside cells
                hdr_paragraphs = [Paragraph(f"<b>{h}</b>", ParagraphStyle('Hdr', parent=body_style, textColor=colors.white)) for h in headers]
                tbl_data = [hdr_paragraphs]
                
                for row in rows:
                    row_paragraphs = []
                    for cell in row:
                        row_paragraphs.append(Paragraph(str(cell), body_style))
                    tbl_data.append(row_paragraphs)
                
                # Limit width
                col_width = 500 / max(len(headers), 1)
                t = Table(tbl_data, colWidths=[col_width] * len(headers))
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#7C3AED')),
                    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                    ('TOPPADDING', (0,0), (-1,-1), 6),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
                    ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F9FAFB')]),
                ]))
                story.append(t)
                story.append(Spacer(1, 10))

        doc.build(story)


# Factory mapping export formats to Exporter strategy objects
EXPORTERS_MAP: Dict[str, BaseExporter] = {
    "Markdown": MarkdownExporter(),
    "HTML": HTMLExporter(),
    "PDF": PDFExporter(),
    "DOCX": DOCXExporter(),
    "TXT": PlainTextExporter(),
    "JSON": JSONExporter(),
    "YAML": YAMLExporter(),
    "XML": XMLExporter(),
    "DRAWIO": DrawioExporter(),
    "PPTX": PowerPointExporter(),
    "JSON (Swagger)": JSONExporter(),  # Maps to JSON exporter
}


def get_exporter(format_name: str) -> BaseExporter:
    """Returns the exporter strategy for the given format name."""
    fmt = format_name.strip()
    
    # Matching
    for key, exporter in EXPORTERS_MAP.items():
        if key.lower() == fmt.lower():
            return exporter
            
    # Normalize extensions/variations
    if fmt.lower() in (".md", "markdown"):
        return EXPORTERS_MAP["Markdown"]
    elif fmt.lower() in (".html", "html"):
        return EXPORTERS_MAP["HTML"]
    elif fmt.lower() in (".pdf", "pdf", "pdf report"):
        return EXPORTERS_MAP["PDF"]
    elif fmt.lower() in (".docx", "docx", "word", "word (.docx)", "microsoft word (.docx)"):
        return EXPORTERS_MAP["DOCX"]
    elif fmt.lower() in (".txt", "txt", "plain text", "plain text (.txt)"):
        return EXPORTERS_MAP["TXT"]
    elif fmt.lower() in (".json", "json", "json (.json)"):
        return EXPORTERS_MAP["JSON"]
    elif fmt.lower() in (".yaml", ".yml", "yaml", "yaml (.yaml / .yml)"):
        return EXPORTERS_MAP["YAML"]
    elif fmt.lower() in (".xml", "xml", "xml (.xml)"):
        return EXPORTERS_MAP["XML"]
    elif fmt.lower() in (".drawio", "drawio", "draw.io diagram (.drawio)", "draw.io"):
        return EXPORTERS_MAP["DRAWIO"]
    elif fmt.lower() in (".pptx", "pptx", "powerpoint", "powerpoint (.pptx)", "microsoft powerpoint (.pptx)"):
        return EXPORTERS_MAP["PPTX"]
        
    raise ValueError(f"Unsupported export format: {format_name}")
